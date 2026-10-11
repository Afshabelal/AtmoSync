
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

NAVY = RGBColor(20, 40, 75)
TEAL = RGBColor(0, 150, 150)
WHITE = RGBColor(255, 255, 255)
DARK = RGBColor(45, 55, 70)

slides_data = [
    ("AtmoSync", [
        "Micro-Climate Arbitrage Analytics",
        "Internship Project | Infotact Solutions",
        "Role: Data Associate L1 Intern",
        "Domain: Data Analytics"
    ]),
    ("Project Overview & Objectives", [
        "Analyze hourly weather patterns across 6 Bengaluru locations.",
        "Collect, inspect and clean weather data.",
        "Perform feature engineering and exploratory data analysis.",
        "Compare micro-climate conditions and identify potential opportunities.",
        "Create visualizations and an interactive dashboard."
    ]),
    ("Dataset & Technology Stack", [
        "Total records: 52,560 hourly observations",
        "Locations: 6 | Records per location: 8,760",
        "Locations: Bengaluru Central, Electronic City, Jayanagar,",
        "Rajajinagar, Whitefield and Yelahanka",
        "Tools: Python, Pandas, NumPy, Matplotlib, Seaborn, Plotly, Streamlit"
    ]),
    ("Project Workflow", [
        "1. Data collection",
        "2. Data inspection",
        "3. Data cleaning and validation",
        "4. Feature engineering",
        "5. Exploratory data analysis",
        "6. Visualization and opportunity screening",
        "7. Dashboard and final reporting"
    ]),
    ("Data Quality & Feature Engineering", [
        "Missing values after cleaning: 0",
        "Duplicate rows after cleaning: 0",
        "Opportunity-flagged records: 2",
        "Derived date, hour, day and month features.",
        "Created temperature/humidity categories and comparison metrics.",
        "Generated an opportunity score and flag."
    ]),
    ("Location-wise Analysis", [
        "Electronic City: lowest average temperature (23.06°C).",
        "Electronic City: highest average humidity (70.66%).",
        "Bengaluru Central and Rajajinagar: highest average temperature (23.76°C).",
        "These two locations also had the lowest average humidity (66.55%).",
        "Insert charts: avg_temp_by_location.png",
        "Insert chart: average_humidity_by_location.png"
    ]),
    ("Monthly Weather Trends", [
        "Warmest month: April — 28.34°C average temperature.",
        "Coolest month: December — 20.09°C average temperature.",
        "Highest average humidity: October — 81.16%.",
        "Lowest average humidity: April — 45.20%.",
        "Insert chart: monthly_temperature_trend.png"
    ]),
    ("Opportunity Score & Findings", [
        "Score combines temperature (up to 70 points) and humidity (up to 30 points).",
        "Screening threshold: score ≥ 60.",
        "18 Mar 2026, 16:00 — Yelahanka: score 84.26.",
        "18 May 2026, 16:00 — Yelahanka: score 100.00.",
        "These are analytical screening signals, not proof of financial profitability."
    ]),
    ("Limitations & Future Scope", [
        "Analysis depends on the quality and representativeness of the dataset.",
        "The opportunity score is a heuristic, not a validated financial model.",
        "Future scope: integrate live weather APIs.",
        "Add forecasting, automated alerts and broader validation.",
        "Evaluate opportunities using real-world cost and demand data."
    ]),
    ("Conclusion & Thank You", [
        "Processed and analyzed 52,560 weather records across 6 locations.",
        "Completed cleaning, feature engineering, analysis and visualization.",
        "Identified seasonal and location-level weather differences.",
        "Built a foundation for further micro-climate opportunity analysis.",
        "Thank You!"
    ])
]

for index, (heading, bullets) in enumerate(slides_data):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = WHITE

    # Header band
    header = slide.shapes.add_shape(
        1, Inches(0), Inches(0), Inches(13.333), Inches(1.25)
    )
    header.fill.solid()
    header.fill.fore_color.rgb = NAVY
    header.line.fill.background()

    title_box = slide.shapes.add_textbox(
        Inches(0.65), Inches(0.25), Inches(12), Inches(0.7)
    )
    title_tf = title_box.text_frame
    title_tf.clear()
    p = title_tf.paragraphs[0]
    p.text = heading
    p.font.name = "Aptos Display"
    p.font.size = Pt(30 if index else 36)
    p.font.bold = True
    p.font.color.rgb = WHITE

    if index == 0:
        start_y = 2.0
    else:
        start_y = 1.65

    body = slide.shapes.add_textbox(
        Inches(0.9), Inches(start_y), Inches(11.6), Inches(4.9)
    )
    tf = body.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.15)
    tf.margin_right = Inches(0.15)

    for j, item in enumerate(bullets):
        p = tf.paragraphs[0] if j == 0 else tf.add_paragraph()
        p.text = item if index == 0 else "•  " + item
        p.font.name = "Aptos"
        p.font.size = Pt(25 if index == 0 else 21)
        p.font.color.rgb = TEAL if index == 0 else DARK
        p.space_after = Pt(18)
        if index == 0:
            p.alignment = PP_ALIGN.CENTER

    # Footer and slide number
    footer = slide.shapes.add_textbox(
        Inches(0.7), Inches(7.05), Inches(12), Inches(0.25)
    )
    fp = footer.text_frame.paragraphs[0]
    fp.text = f"AtmoSync | Infotact Solutions                                      {index + 1} / {len(slides_data)}"
    fp.font.size = Pt(10)
    fp.font.color.rgb = TEAL

output_path = "AtmoSync_Presentation.pptx"
prs.save(output_path)
print("Presentation created successfully!")
print(f"Slides: {len(prs.slides)}")
print(f"File: {output_path}")
